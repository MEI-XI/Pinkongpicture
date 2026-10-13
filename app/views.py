from pathlib import Path
from uuid import uuid4
from zipfile import BadZipFile

from fastapi import APIRouter, Request, UploadFile
from fastapi.responses import FileResponse
from fastapi.templating import Jinja2Templates
from openpyxl import load_workbook
from pydantic import BaseModel, Field

from app.domain import FIELDS, FieldValue
from app.excel import ExcelExporter, TemplateMapping
from app.extraction import extract_candidates
from app.grouping import propose_groups
from app.samples import Samples

router = APIRouter()
templates = Jinja2Templates(directory=Path(__file__).parent / "templates")


class ImageSelection(BaseModel):
    image_ids: list[str]


class RecordInput(BaseModel):
    fields: dict[str, FieldValue]
    operator: str
    expected_version: int
    grouping_reviewed: bool = False
    reason: str = ""
    product_decision: dict = Field(default_factory=lambda: {"mode": "new"})
    region_labels: dict[str, str] = Field(default_factory=dict)
    reviewed_region_refs: list[str] = Field(default_factory=list)


class VoidInput(BaseModel):
    operator: str
    reason: str
    expected_version: int


class MoveInput(BaseModel):
    target_id: str | None
    operator: str
    reason: str
    expected_versions: dict[str, int]


@router.get("/")
def workbench(request: Request):
    return templates.TemplateResponse(request=request, name="workbench.html", context={"fields": FIELDS})


@router.get("/api/state")
def state(request: Request):
    records = request.app.state.records
    settings = records.store.list("settings")
    config = next((value for value in settings if value.get("id") == "template"), None)
    person = next((value for value in settings if value.get("id") == "operator"), {})
    return {
        "images": [
            image.model_dump(mode="json", exclude={"path"}) for image in request.app.state.images.list()
        ],
        "receipts": records.list(),
        "products": records.products.list(),
        "samples": Samples(request.app.state.root).list(),
        "operator": person.get("name", ""),
        "template": config,
    }


@router.post("/api/images")
def import_images(request: Request, files: list[UploadFile]):
    imported, errors = [], []
    batch_id = uuid4().hex
    for file in files:
        try:
            imported.append(
                request.app.state.images.import_image(
                    file.filename or "图片", file.file.read(30 * 1024 * 1024 + 1), import_batch=batch_id
                )
            )
        except ValueError as exc:
            errors.append({"name": file.filename, "error": str(exc)})
        finally:
            file.file.close()
    request.app.state.ocr.submit([image.id for image in imported])
    return {
        "images": [image.model_dump(mode="json", exclude={"path"}) for image in imported],
        "errors": errors,
    }


@router.get("/api/images/{image_id}/file")
def image_file(request: Request, image_id: str):
    return FileResponse(request.app.state.images.get(image_id).path, media_type="image/png")


@router.get("/api/samples/{sample_id}/file")
def sample_file(request: Request, sample_id: str):
    sample = request.app.state.records.store.get("sample", sample_id)
    path = Path(sample["crop_path"])
    if not path.is_file():
        raise KeyError(sample_id)
    return FileResponse(path, media_type="image/png")


@router.post("/api/receipts")
def create_receipt(request: Request, body: ImageSelection):
    return request.app.state.records.create_receipt(body.image_ids)


@router.post("/api/images/{image_id}/retry")
def retry_image(request: Request, image_id: str):
    request.app.state.ocr.submit([image_id])
    return {"queued": True}


@router.put("/api/receipts/{receipt_id}")
def edit(request: Request, receipt_id: str, body: RecordInput):
    return request.app.state.records.edit_receipt(receipt_id, **body.model_dump())


@router.post("/api/receipts/{receipt_id}/void")
def void_receipt(request: Request, receipt_id: str, body: VoidInput):
    return request.app.state.records.void_receipt(receipt_id, **body.model_dump())


@router.post("/api/images/{image_id}/move")
def move_image(request: Request, image_id: str, body: MoveInput):
    request.app.state.records.move_image(image_id, **body.model_dump())
    return request.app.state.records.list()


@router.post("/api/groups/suggest")
def group_suggestions(request: Request, body: ImageSelection):
    assigned = {
        image for receipt in request.app.state.records.list() for image in receipt["draft"]["image_ids"]
    }
    if assigned.intersection(body.image_ids):
        raise ValueError("请选择待分组照片，已归组照片请使用移动操作")
    batches = {}
    for identifier in body.image_ids:
        asset = request.app.state.images.get(identifier)
        batches.setdefault(asset.import_batch, {})[identifier] = extract_candidates(asset.blocks)
    combined = {"groups": [], "ungrouped": []}
    for candidates in batches.values():
        proposal = propose_groups(candidates)
        for key in combined:
            combined[key].extend(proposal[key])
    return combined


@router.post("/api/products/suggest")
def product_suggestions(request: Request, body: dict[str, str]):
    return request.app.state.records.products.suggest_products(body)


@router.get("/api/receipts/{receipt_id}/candidates")
def candidates(request: Request, receipt_id: str):
    receipt = request.app.state.records.get(receipt_id)
    blocks = [
        block
        for image_id in receipt["draft"]["image_ids"]
        for block in request.app.state.images.get(image_id).blocks
    ]
    return extract_candidates(blocks)


@router.post("/api/receipts/{receipt_id}/confirm")
def confirm(request: Request, receipt_id: str, body: RecordInput):
    records = request.app.state.records
    arguments = body.model_dump()
    record = records.confirm_receipt(receipt_id, **arguments)
    records.store.put("settings", "operator", {"id": "operator", "name": body.operator})
    try:
        samples = Samples(request.app.state.root).collect_samples(record, record["region_labels"])
        sample_status = {"status": "saved", "count": len(samples), "error": ""}
    except (OSError, ValueError) as exc:
        sample_status = {"status": "error", "count": 0, "error": str(exc)}
    records.store.put("sample_status", f"{receipt_id}:{record['version']}", sample_status)
    return {"record": record, "samples": sample_status}


@router.post("/api/template")
def upload_template(request: Request, file: UploadFile):
    if not file.filename or Path(file.filename).suffix.lower() != ".xlsx":
        raise ValueError("请上传 .xlsx 模板，旧版 .xls 需先转换")
    directory = request.app.state.root / "templates"
    directory.mkdir(exist_ok=True)
    path = directory / f"{uuid4().hex}.xlsx"
    content = file.file.read(10 * 1024 * 1024 + 1)
    file.file.close()
    if len(content) > 10 * 1024 * 1024:
        raise ValueError("模板不能超过 10 MB")
    path.write_bytes(content)
    try:
        book = load_workbook(path)
    except (BadZipFile, OSError) as exc:
        path.unlink(missing_ok=True)
        raise ValueError("无法读取模板，请检查文件格式") from exc
    config = {
        "id": "template",
        "path": str(path),
        "name": file.filename,
        "sheets": book.sheetnames,
        "config": None,
    }
    book.close()
    request.app.state.records.store.put("settings", "template", config)
    return {"sheets": config["sheets"], "name": config["name"]}


@router.put("/api/template/config")
def template_config(request: Request, body: dict):
    store = request.app.state.records.store
    config = store.get("settings", "template")
    mapping = TemplateMapping.model_validate(body)
    if mapping.sheet not in config["sheets"]:
        raise ValueError("模板中没有选择的工作表")
    config["config"] = mapping.model_dump()
    store.put("settings", "template", config)
    return {"saved": True}


@router.post("/api/exports")
def export(request: Request, body: dict):
    config = request.app.state.records.store.get("settings", "template")
    if config["config"] is None:
        raise ValueError("请先配置模板的工作表、数据区域和字段列")
    result = ExcelExporter(request.app.state.root).export_confirmed(
        Path(config["path"]), config["config"], body
    )
    return {
        "count": result["count"],
        "pending_count": result["pending_count"],
        "url": f"/api/exports/{result['path'].name}",
        "name": result["path"].name,
    }


@router.get("/api/exports/{filename}")
def download(request: Request, filename: str):
    if (
        Path(filename).name != filename
        or not filename.startswith("来料记录-")
        or not filename.endswith(".xlsx")
    ):
        raise ValueError("导出文件名无效")
    path = request.app.state.root / "exports" / filename
    if not path.is_file():
        raise KeyError(filename)
    return FileResponse(
        path,
        filename=filename,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
