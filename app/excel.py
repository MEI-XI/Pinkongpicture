from copy import copy
from datetime import date
from pathlib import Path
from typing import Literal
from uuid import uuid4

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.formula.translate import Translator
from openpyxl.utils import column_index_from_string
from pydantic import BaseModel, Field, model_validator

from app.domain import FIELDS
from app.records import Records


class TemplateMapping(BaseModel):
    sheet: str = Field(min_length=1)
    start_row: int = Field(ge=1, le=1048576)
    end_row: int = Field(ge=1, le=1048576)
    columns: dict[str, str]
    date_mode: Literal["raw", "normalized"]
    missing_mode: Literal["with_reason", "label"]

    @model_validator(mode="after")
    def validate_region(self):
        if self.end_row < self.start_row:
            raise ValueError("数据结束行不能早于起始行")
        if set(self.columns) != {*FIELDS, "entered_at", "operator"}:
            raise ValueError("请映射七个字段、录入时间和操作人")
        self.columns = {key: value.strip().upper() for key, value in self.columns.items()}
        numbers = [column_index_from_string(value) for value in self.columns.values()]
        if len(set(numbers)) != 9 or any(number > 16384 for number in numbers):
            raise ValueError("每个字段必须映射到不同的有效列")
        return self


class ExcelExporter:
    def __init__(self, root: Path):
        self.records = Records(root)
        self.directory = Path(root) / "exports"
        self.directory.mkdir(parents=True, exist_ok=True)

    def export_confirmed(self, template: Path, config: dict, selection=None):
        config = TemplateMapping.model_validate(config).model_dump()
        selection = selection or {}
        for key in ("date_from", "date_to"):
            if selection.get(key):
                date.fromisoformat(selection[key])
        if (
            selection.get("date_from")
            and selection.get("date_to")
            and selection["date_from"] > selection["date_to"]
        ):
            raise ValueError("起始日期不能晚于结束日期")
        if Path(template).suffix.lower() != ".xlsx":
            raise ValueError("当前仅支持 .xlsx 模板，请先转换旧版 .xls 并核对格式")
        receipts = self.records.list()
        rows = [
            receipt["confirmed"] for receipt in receipts if receipt["confirmed"] and not receipt.get("voided")
        ]
        rows = [
            row
            for row in rows
            if (
                (not selection.get("date_from") or row["entered_at"][:10] >= selection["date_from"])
                and (not selection.get("date_to") or row["entered_at"][:10] <= selection["date_to"])
                and (not selection.get("lot") or row["fields"]["lot"]["value"] == selection["lot"])
                and (
                    not selection.get("product_id")
                    or (row.get("product") or {}).get("id") == selection["product_id"]
                )
            )
        ]
        book = load_workbook(template)
        if config["sheet"] not in book.sheetnames:
            raise ValueError("模板中没有选择的工作表")
        sheet = book[config["sheet"]]
        start = config["start_row"]
        end = config["end_row"]
        if not 1 <= start <= end <= 1048576 or len(rows) > end - start + 1:
            raise ValueError("记录数超过模板数据区域，或数据行范围无效")
        if set(config["columns"]) != {*FIELDS, "entered_at", "operator"}:
            raise ValueError("请映射七个字段、录入时间和操作人")
        columns = {
            field: column_index_from_string(column.upper()) for field, column in config["columns"].items()
        }
        if len(set(columns.values())) != 9 or any(column > 16384 for column in columns.values()):
            raise ValueError("每个字段必须映射到不同的有效列")
        if config["date_mode"] not in {"raw", "normalized"} or config["missing_mode"] not in {
            "with_reason",
            "label",
        }:
            raise ValueError("请选择日期和未标注信息的输出方式")
        original_last_row = sheet.max_row
        for row in range(start, min(end, max(sheet.max_row, start)) + 1):
            for column in columns.values():
                cell = sheet.cell(row, column)
                if isinstance(cell, MergedCell) or cell.value is not None:
                    raise ValueError(f"{cell.coordinate} 不是空白数据单元格，请检查映射或清空模板示例值")
        for offset, record in enumerate(rows):
            row = start + offset
            if offset:
                if row > original_last_row:
                    sheet.row_dimensions[row].height = sheet.row_dimensions[start].height
                for column in range(1, sheet.max_column + 1):
                    source, target = sheet.cell(start, column), sheet.cell(row, column)
                    if row > original_last_row:
                        target._style = copy(source._style)
                    if source.data_type == "f" and target.value is None and column not in columns.values():
                        target.value = Translator(source.value, origin=source.coordinate).translate_formula(
                            target.coordinate
                        )
            for field, column in columns.items():
                if field in {"entered_at", "operator"}:
                    value = record[field]
                else:
                    item = record["fields"][field]
                    value = (
                        item["raw"]
                        if field == "production_date" and config["date_mode"] == "raw"
                        else item["value"]
                    )
                    if item["source"] == "missing":
                        value = "未标注" + (
                            f"（{item['note']}）" if config["missing_mode"] == "with_reason" else ""
                        )
                cell = sheet.cell(row, column)
                cell.value = value
                cell.data_type = "s"
        path = self.directory / f"来料记录-{uuid4().hex}.xlsx"
        book.save(path)
        selected_ids = {row["id"] for row in rows}
        return {
            "path": path,
            "count": len(rows),
            "pending_count": sum(
                bool(receipt["id"] in selected_ids and receipt["pending"]) for receipt in receipts
            ),
        }
