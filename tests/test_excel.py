from datetime import datetime, timezone

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

from tests.test_records import make_image, verified_fields


def template(root):
    path = root / "provided-template.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.title = "来料"
    sheet["A1"] = "产品名称"
    sheet["E2"].number_format = "@"
    sheet["A2"].font = Font(name="微软雅黑", bold=True)
    sheet["A2"].fill = PatternFill("solid", fgColor="DDEEFF")
    sheet["J2"] = "=ROW()-1"
    book.save(path)
    return path


def mapping():
    return {
        "sheet": "来料",
        "start_row": 2,
        "end_row": 20,
        "columns": dict(
            zip(
                [
                    "name",
                    "specification",
                    "manufacturer",
                    "address",
                    "lot",
                    "production_date",
                    "shelf_life",
                    "entered_at",
                    "operator",
                ],
                "ABCDEFGHI",
            )
        ),
        "date_mode": "raw",
        "missing_mode": "with_reason",
    }


def test_exports_only_confirmed_snapshot_and_preserves_template_and_leading_zeros(tmp_path):
    from app.excel import ExcelExporter
    from app.records import Records

    records = Records(tmp_path / "data")
    receipt = records.create_receipt([make_image(tmp_path / "data")])
    records.confirm_receipt(
        receipt["id"], fields=verified_fields(), operator="王检验", expected_version=0, grouping_reviewed=True
    )
    records.create_receipt([make_image(tmp_path / "data")])
    records.edit_receipt(
        receipt["id"],
        fields=verified_fields(lot="00555"),
        operator="王检验",
        expected_version=1,
        reason="更正中",
    )
    source = template(tmp_path)
    original = source.read_bytes()
    exporter = ExcelExporter(tmp_path / "data")
    output = exporter.export_confirmed(source, mapping())
    sheet = load_workbook(output["path"]).active
    assert output["count"] == 1
    assert output["pending_count"] == 1
    assert sheet["E2"].value == "00123"
    assert sheet["E2"].data_type == "s"
    assert sheet["A2"].font.bold is True
    assert sheet["A2"].fill.fgColor.rgb == "00DDEEFF"
    assert sheet["A3"].value is None
    assert sheet["J2"].value == "=ROW()-1"
    assert source.read_bytes() == original
    saved = output["path"].read_bytes()
    second = exporter.export_confirmed(source, mapping())
    assert second["path"] != output["path"]
    assert output["path"].read_bytes() == saved


def test_export_filter_uses_original_entry_date_and_writes_untrusted_text_as_text(tmp_path):
    from app.excel import ExcelExporter
    from app.records import Records

    now = [datetime(2026, 10, 13, 1, tzinfo=timezone.utc)]
    records = Records(tmp_path / "data", clock=lambda: now[0])
    fields = verified_fields(name="=DANGEROUS()")
    fields["shelf_life"] = {"source": "missing", "note": "包装未列保质期"}
    first = records.create_receipt([make_image(tmp_path / "data")])
    records.confirm_receipt(
        first["id"], fields=fields, operator="王检验", expected_version=0, grouping_reviewed=True
    )
    now[0] = datetime(2026, 10, 14, 1, tzinfo=timezone.utc)
    second = records.create_receipt([make_image(tmp_path / "data")])
    records.confirm_receipt(
        second["id"],
        fields=verified_fields(lot="00009"),
        operator="王检验",
        expected_version=0,
        grouping_reviewed=True,
    )
    export = ExcelExporter(tmp_path / "data").export_confirmed(
        template(tmp_path), mapping(), {"date_from": "2026-10-13", "date_to": "2026-10-13", "lot": "00123"}
    )
    sheet = load_workbook(export["path"]).active
    assert export["count"] == 1
    assert sheet["A2"].value == "=DANGEROUS()"
    assert sheet["A2"].data_type == "s"
    assert sheet["G2"].value == "未标注（包装未列保质期）"
    assert sheet["F2"].value == "2026年10月1日"


def test_template_mapping_rejects_collisions_and_does_not_overwrite_existing_cells(tmp_path):
    from app.excel import ExcelExporter

    source = template(tmp_path)
    exporter = ExcelExporter(tmp_path / "data")
    config = mapping()
    config["columns"]["operator"] = "A"
    with pytest.raises(ValueError):
        exporter.export_confirmed(source, config)
    book = load_workbook(source)
    book.active["E2"] = "模板示例，必须先清空"
    book.save(source)
    with pytest.raises(ValueError):
        exporter.export_confirmed(source, mapping())


def test_existing_row_format_is_preserved_and_only_new_rows_inherit_template_style(tmp_path):
    from app.excel import ExcelExporter
    from app.records import Records

    records = Records(tmp_path / "data")
    for lot in ["001", "002", "003"]:
        receipt = records.create_receipt([make_image(tmp_path / "data")])
        records.confirm_receipt(
            receipt["id"],
            fields=verified_fields(lot=lot),
            operator="王检验",
            expected_version=0,
            grouping_reviewed=True,
        )
    source = template(tmp_path)
    book = load_workbook(source)
    book.active.row_dimensions[3].height = 48
    book.active["A3"].font = Font(italic=True)
    book.active["J3"].number_format = "0.00"
    book.save(source)
    exported = ExcelExporter(tmp_path / "data").export_confirmed(source, mapping())
    sheet = load_workbook(exported["path"]).active
    assert sheet.row_dimensions[3].height == 48
    assert sheet["A3"].font.italic is True
    assert sheet["J3"].number_format == "0.00"
    assert sheet["A4"].font.bold is True
    assert sheet["J4"].value == "=ROW()-1"


def test_explicit_void_of_empty_confirmed_group_excludes_duplicate_from_export_with_audit(tmp_path):
    from app.excel import ExcelExporter
    from app.records import Records

    records = Records(tmp_path / "data")
    image_ids = [make_image(tmp_path / "data") for _ in range(2)]
    receipts = [records.create_receipt([identifier]) for identifier in image_ids]
    for receipt in receipts:
        records.confirm_receipt(
            receipt["id"],
            fields=verified_fields(),
            operator="王检验",
            expected_version=0,
            grouping_reviewed=True,
        )
    first, second = receipts
    records.move_image(
        image_ids[0],
        second["id"],
        operator="李检验",
        reason="核实为同一次来料，误拆成两组",
        expected_versions={first["id"]: 1, second["id"]: 1},
    )
    assert len(records.list_confirmed()) == 2  # Moving alone must not silently void a confirmation.
    with pytest.raises(ValueError):
        records.void_receipt(first["id"], operator="李检验", reason="", expected_version=2)
    with pytest.raises(ValueError):
        records.void_receipt(
            second["id"], operator="李检验", reason="仍有图片，不得误作废", expected_version=2
        )
    records.void_receipt(
        first["id"], operator="李检验", reason="重复登记，照片已归入另一条记录", expected_version=2
    )
    assert len(records.list_confirmed()) == 1
    preserved = Records(tmp_path / "data").get(first["id"])
    assert preserved["voided"]["operator"] == "李检验"
    assert preserved["voided"]["reason"] == "重复登记，照片已归入另一条记录"
    assert len(preserved["history"]) == 1
    assert preserved["confirmed"]["fields"]["lot"]["value"] == "00123"
    with pytest.raises(ValueError):
        records.move_image(
            image_ids[0],
            first["id"],
            operator="李检验",
            reason="不允许归入作废记录",
            expected_versions={first["id"]: 3, second["id"]: 2},
        )
    output = ExcelExporter(tmp_path / "data").export_confirmed(template(tmp_path), mapping())
    assert output["count"] == 1
